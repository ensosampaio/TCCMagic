import json
from collections import Counter
from typing import Sequence

import numpy as np
import pytest

from tccmagic.analysis import analyze_swaps, metagame_effects
from tccmagic.archive import GameArchive
from tccmagic.brkga import BRKGA_Optimizer, BRKGAConfig
from tccmagic.cards import Deck, Opponent, load_database, load_reference_sideboard
from tccmagic.decoder import SideboardDecoder
from tccmagic.export import export_comparison_json, export_comparison_xlsx, export_json, export_xlsx
from tccmagic.fitness import FitnessEvaluator
from tccmagic.pipeline import (
    NO_SIDEBOARD,
    Candidate,
    ExperimentConfig,
    Validation,
    compare_sideboards,
    measure_game1,
    paired_comparison,
    random_sideboard,
    run_experiment,
    sideboard_from_names,
)
from tccmagic.simulation import Bo3Simulator, ForgeConfig, ForgeError, GameSpec, MatchEngine, SurrogateEngine
from tccmagic.simulation.forge import OPPONENT_TAG, PLAYER_TAG, parse_forge_output


@pytest.fixture(scope="module")
def db():
    return load_database()


class ScriptedEngine(MatchEngine):
    """Motor falso: resultado definido apenas pelo número do jogo."""

    def __init__(self, outcome_by_game: dict[int, bool]):
        self.outcome_by_game = outcome_by_game
        self.calls: list[list[GameSpec]] = []

    def play_games(self, deck: Deck, opponent: Opponent, specs: Sequence[GameSpec]) -> list[bool]:
        self.calls.append(list(specs))
        return [self.outcome_by_game[s.game_number] for s in specs]


class CoinEngine(MatchEngine):
    """Motor ruidoso em que o deck não importa: todo jogo é uma moeda viciada."""

    def __init__(self, p: float = 0.6, seed: int = 0):
        self.p = p
        self.rng = np.random.default_rng(seed)

    def play_games(self, deck: Deck, opponent: Opponent, specs: Sequence[GameSpec]) -> list[bool]:
        return [bool(u < self.p) for u in self.rng.random(len(specs))]


# --------------------------------------------------------------------------- #
# Orquestração Bo3
# --------------------------------------------------------------------------- #

def test_bo3_plays_game3_only_when_tied_and_scores_matches(db):
    engine = ScriptedEngine({1: True, 2: False, 3: True})
    sim = Bo3Simulator(db, engine, matches_per_opponent=4)
    report = sim.evaluate(())
    for m in report.matchups:
        assert m.winrate_md1 == 1.0 and m.winrate_bo3 == 1.0 and m.winrate_postboard == 0.5
    # G3 é jogado em todas as séries, pois todas ficam 1–1; o perdedor do G2 (nós) começa.
    g3_calls = [c for c in engine.calls if c and c[0].game_number == 3]
    assert all(s.on_play for c in g3_calls for s in c)


def test_bo3_skips_game3_after_two_wins(db):
    engine = ScriptedEngine({1: True, 2: True, 3: False})
    report = Bo3Simulator(db, engine, matches_per_opponent=3).evaluate(())
    assert report.winrate_bo3 == 1.0
    assert not any(c and c[0].game_number == 3 for c in engine.calls)


def test_game1_alternates_play_draw_and_is_computed_once(db):
    engine = ScriptedEngine({1: False, 2: False, 3: False})
    sim = Bo3Simulator(db, engine, matches_per_opponent=4)
    sim.evaluate(())
    sim.evaluate(())
    g1_calls = [c for c in engine.calls if c and c[0].game_number == 1]
    assert len(g1_calls) == len(db.opponents)
    assert [s.on_play for s in g1_calls[0]] == [True, False, True, False]


# --------------------------------------------------------------------------- #
# Plano de troca aleatório e cache de jogos
# --------------------------------------------------------------------------- #

def full_pool(db):
    """As 15 primeiras entradas do pool expandido (um sideboard legal qualquer)."""
    return [slot.card for slot in db.sideboard_slots()][:15]


def test_random_plan_draws_from_sideboard_and_keeps_lands(db):
    sim = Bo3Simulator(db, CoinEngine(), plan_block=1, max_swaps=5)
    side = full_pool(db)
    lands = sum(n for c, n in db.maindeck.cards if c.is_land)
    sizes = Counter()
    for opp in db.opponents:
        for m in range(200):
            plan = sim.plan_for(opp, side, m)
            sizes[plan.n_swaps] += 1
            assert not Counter(c.name for c in plan.cards_in) - Counter(c.name for c in side)
            assert all(c.flex and not c.is_land for c in plan.cards_out)
            assert not {c.name for c in plan.cards_in} & {c.name for c in plan.cards_out}
            assert plan.deck.size == 60
            assert sum(n for c, n in plan.deck.cards if c.is_land) == lands
            assert sim.plan_for(opp, list(reversed(side)), m) is plan  # ordem do side não importa
    assert set(sizes) == {1, 2, 3, 4, 5}


def test_random_plan_is_uniform_over_sideboard_copies(db):
    sim = Bo3Simulator(db, CoinEngine(), plan_block=1, max_swaps=5)
    side = full_pool(db)
    opp = db.opponents[0]
    n = 3000
    drawn = Counter(c.name for m in range(n) for c in sim.plan_for(opp, side, m).cards_in)
    # Nº médio de trocas = 3; cada uma das 15 cópias entra em ~3/15 das séries.
    for name, copies in Counter(c.name for c in side).items():
        assert drawn[name] / n == pytest.approx(copies * 3 / 15, abs=0.04)


def test_random_plan_is_shared_by_sideboards_holding_the_drawn_cards(db):
    """Números aleatórios comuns: trocar uma carta não sorteada não muda o plano."""
    sim = Bo3Simulator(db, CoinEngine(), plan_block=1, max_swaps=5)
    side = full_pool(db)
    spare = [slot.card for slot in db.sideboard_slots()][15:]
    opp = db.opponents[1]
    shared = 0
    for m in range(100):
        plan = sim.plan_for(opp, side, m)
        drawn = {c.name for c in plan.cards_in}
        idle = next((i for i, c in enumerate(side) if c.name not in drawn and c.name != spare[0].name), None)
        if idle is None:
            continue
        other = side[:idle] + [spare[0]] + side[idle + 1:]
        other_plan = sim.plan_for(opp, other, m)
        if spare[0].name not in {c.name for c in other_plan.cards_in}:
            assert other_plan.signature == plan.signature
            shared += 1
    assert shared > 30


def test_game_cache_simulates_each_deck_once(db):
    engine = ScriptedEngine({1: True, 2: False, 3: True})
    sim = Bo3Simulator(db, engine, matches_per_opponent=20, plan_block=5)
    side = full_pool(db)
    sim.prepare()
    first, second = sim.play_postboard([(side, 0), (list(reversed(side)), 0)], sim.game1)
    assert first == second
    assert sim.games_reused == sim.games_played  # o 2º sideboard gera os mesmos decks
    calls = len(engine.calls)
    sim.evaluate(side)  # mesmas séries de novo: nada é simulado
    assert len(engine.calls) == calls
    # Um lote por (oponente, deck): cada bloco de 5 séries tem o seu sorteio.
    g2_calls = [c for c in engine.calls if c and c[0].game_number == 2]
    assert len(g2_calls) <= len(db.opponents) * 4


def test_affinity_plan_is_the_same_in_every_series(db):
    sim = Bo3Simulator(db, ScriptedEngine({1: True, 2: True, 3: True}), matches_per_opponent=7, plan="afinidade")
    report = sim.evaluate(full_pool(db))
    for m in report.matchups:
        assert all(total % 7 == 0 for _, total in m.swaps_in)
        assert "." not in m.cards_in  # cópias inteiras por série
    with pytest.raises(ValueError):
        Bo3Simulator(db, CoinEngine(), plan="outro")


class WearTearEngine(MatchEngine):
    """Cada cópia de Wear // Tear no deck soma 15 pp à chance de vitória; nada mais importa."""

    def __init__(self, seed: int = 0):
        self.rng = np.random.default_rng(seed)

    def play_games(self, deck: Deck, opponent: Opponent, specs: Sequence[GameSpec]) -> list[bool]:
        p = 0.35 + 0.15 * deck.counts().get("Wear // Tear", 0)
        return [bool(u < p) for u in self.rng.random(len(specs))]


def test_swap_analysis_recovers_the_effect_of_a_card(db):
    sim = Bo3Simulator(db, WearTearEngine(), matches_per_opponent=1500, plan_block=1)
    sim.evaluate(full_pool(db))
    analyses = analyze_swaps(sim.games, db)
    assert [a.opponent for a in analyses] == [o.name for o in db.opponents]
    for a in analyses:
        effects = {e.card: e for e in a.role("entra")}
        wear = effects.pop("Wear // Tear")
        assert wear.effect == pytest.approx(0.15, abs=0.05) and wear.significant
        assert all(abs(e.effect) < 0.08 for e in effects.values())
        assert a.no_swap_winrate == pytest.approx(0.35, abs=0.06)
    assert metagame_effects(analyses)[0].card == "Wear // Tear"


# --------------------------------------------------------------------------- #
# Simulador substituto
# --------------------------------------------------------------------------- #

def test_surrogate_game1_probability_depends_only_on_base_logit(db):
    engine = SurrogateEngine(db.maindeck)
    engine.config.play_bonus = 0.0
    for opp in db.opponents:
        p = engine.win_probability(db.maindeck, opp, game_number=1, on_play=True)
        assert p == pytest.approx(1 / (1 + np.exp(-opp.base_logit)))


def test_surrogate_rewards_enchantment_hate_against_bogles(db):
    engine = SurrogateEngine(db.maindeck)
    bogles = next(o for o in db.opponents if o.name == "Bogles")
    wear_tear = next(c for c in db.candidates if c.name == "Wear // Tear")
    discharge = next(c for c, _ in db.maindeck.cards if c.name == "Galvanic Discharge")
    boarded = db.maindeck.swap([discharge] * 3, [wear_tear] * 3)
    assert engine.win_probability(boarded, bogles, 2, True) > engine.win_probability(
        db.maindeck, bogles, 2, True
    )


def test_surrogate_uses_common_random_numbers(db):
    a = Bo3Simulator(db, SurrogateEngine(db.maindeck), matches_per_opponent=30).evaluate(())
    b = Bo3Simulator(db, SurrogateEngine(db.maindeck), matches_per_opponent=30).evaluate(())
    assert a == b


# --------------------------------------------------------------------------- #
# Fitness: cache e paralelismo
# --------------------------------------------------------------------------- #

def test_fitness_cache_and_parallel_evaluation_match_serial(db):
    decoder = SideboardDecoder(db)
    population = np.random.default_rng(3).random((12, 9))

    serial = FitnessEvaluator(decoder, Bo3Simulator(db, SurrogateEngine(db.maindeck), 20))
    f1 = serial(population)
    simulated = serial.simulations
    np.testing.assert_array_equal(serial(population), f1)
    assert serial.simulations == simulated  # tudo veio do cache

    with FitnessEvaluator(decoder, Bo3Simulator(db, SurrogateEngine(db.maindeck), 20), workers=2) as par:
        np.testing.assert_array_equal(par(population), f1)
        sideboards = [decoder.decode(population[0]).cards, ()]
        assert par.validate(sideboards, 30) == serial.validate(sideboards, 30)


def test_resampling_accumulates_series_per_sideboard(db):
    decoder = SideboardDecoder(db)
    population = np.random.default_rng(3).random((6, 9))
    evaluator = FitnessEvaluator(decoder, Bo3Simulator(db, CoinEngine(), 10), resample=True)

    first = evaluator.evaluate_population(population)
    distinct = len({e.decoded.key for e in first})
    assert all(e.report.n_matches == 10 for e in first) and evaluator.simulations == distinct

    second = evaluator.evaluate_population(population)
    assert all(e.report.n_matches == 20 for e in second) and evaluator.simulations == 2 * distinct
    # O Game 1 é o mesmo em todas as rodadas; só os Games 2/3 são jogados de novo.
    assert [m.game1_wins for m in second[0].report.matchups] == [2 * m.game1_wins for m in first[0].report.matchups]

    # Consultar um indivíduo já avaliado não joga rodadas extras.
    assert evaluator.evaluate_one(population[0]).report.n_matches == 20
    assert evaluator.simulations == 2 * distinct


def test_noise_alone_is_not_reported_as_sideboard_gain(db):
    """Com um motor em que o sideboard não tem efeito, a validação não aponta ganho."""
    decoder = SideboardDecoder(db)
    evaluator = FitnessEvaluator(decoder, Bo3Simulator(db, CoinEngine(), 2), resample=True)
    optimizer = BRKGA_Optimizer(9, evaluator, BRKGAConfig(population_size=10, generations=5, seed=0),
                                reevaluate_elites=True)
    best = evaluator.evaluate_one(optimizer.run().best_keys)
    assert best.report.n_matches >= 4  # o melhor final foi confirmado ao menos uma vez

    validated, baseline = evaluator.validate([best.decoded.cards, ()], 300)
    assert not Validation(300, validated, baseline).significant


# --------------------------------------------------------------------------- #
# Forge (apenas parsing / montagem de comando — não requer Java)
# --------------------------------------------------------------------------- #

FORGE_STDOUT = f"""
Simulation mode
Game 1 ended in 51234 ms. Ai(1)-{PLAYER_TAG} has won!
Game 2 ended in 40211 ms. Ai(2)-{OPPONENT_TAG} has won!
Game 3 ended in a Draw! Took 9999 ms.
"""


def test_parse_forge_output():
    cfg = ForgeConfig(jar_path="forge.jar")
    assert parse_forge_output(FORGE_STDOUT, 3, cfg) == [True, False, False]
    with pytest.raises(ForgeError):
        parse_forge_output(FORGE_STDOUT, 4, cfg)


def test_forge_engine_runs_subprocess_and_parses(db, tmp_path, monkeypatch):
    """Substitui 'java' por um script que imita a saída do Forge."""
    import stat
    import sys

    from tccmagic.simulation import ForgeEngine, forge

    decks_dir = tmp_path / "decks"
    monkeypatch.setattr(forge, "_decks_dir", lambda: decks_dir)  # não toca no perfil real do Forge

    script = tmp_path / "fake_java.py"
    script.write_text(
        "import os, sys\n"
        "args = sys.argv[1:]\n"
        "decks = args[args.index('-d') + 1: args.index('-d') + 3]\n"
        f"decks = [os.path.join({str(decks_dir)!r}, d) for d in decks]\n"
        "assert all(open(d).read().startswith('[metadata]') for d in decks)\n"
        "n = int(args[args.index('-n') + 1])\n"
        "for g in range(1, n + 1):\n"
        f"    who = '{PLAYER_TAG}' if g % 2 else '{OPPONENT_TAG}'\n"
        "    print(f'Game {g} ended in 10 ms. Ai(1)-{who} has won!')\n"
    )
    if sys.platform == "win32":
        fake_java = tmp_path / "fake_java.cmd"
        fake_java.write_text(f'@"{sys.executable}" "{script}" %*\n')
    else:
        fake_java = tmp_path / "fake_java"
        fake_java.write_text(f'#!/bin/sh\nexec "{sys.executable}" "{script}" "$@"\n')
        fake_java.chmod(fake_java.stat().st_mode | stat.S_IEXEC)
    jar = tmp_path / "forge.jar"
    jar.write_text("")

    engine = ForgeEngine(ForgeConfig(jar_path=str(jar), java_executable=str(fake_java)))
    specs = [GameSpec(m, 1, True) for m in range(4)]
    assert engine.play_games(db.maindeck, db.opponents[0], specs) == [True, False, True, False]
    assert not list(decks_dir.iterdir())  # os .dck temporários são removidos


def test_forge_dck_format(db):
    text = db.maindeck.to_forge_dck()
    assert text.startswith("[metadata]\nName=Boros Energy\n[Main]\n")
    assert "4 Ragavan, Nimble Pilferer" in text
    assert sum(int(line.split(" ", 1)[0]) for line in text.splitlines()[3:]) == 60


# --------------------------------------------------------------------------- #
# Pipeline + exportação
# --------------------------------------------------------------------------- #

def test_small_experiment_and_export(tmp_path):
    cfg = ExperimentConfig(matches_per_opponent=20, brkga=BRKGAConfig(population_size=12, generations=3, seed=1))
    result = run_experiment(cfg)
    assert len(result.history) == 4
    assert len(result.best.decoded.slots) == 15
    assert result.best.fitness == max(h.best_fitness for h in result.history)

    data = json.loads(export_json(result, tmp_path / "r.json").read_text(encoding="utf-8"))
    assert len(data["final"]["sideboard"]) == 15 and len(data["history"]) == 4

    from openpyxl import load_workbook

    wb = load_workbook(export_xlsx(result, tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Geracoes", "Pesos_Finais", "Sideboard", "Matchups", "Validacao",
                             "Efeito_Trocas", "Resumo"]
    assert wb["Efeito_Trocas"].max_row > 1
    assert data["swap_analysis"]["opponents"] and data["games_simulated"] > 0
    assert wb["Sideboard"].max_row == 16
    assert wb["Validacao"].max_row == 1 + 2 * len(result.database.opponents)


def test_experiment_validates_best_and_baseline_on_fresh_series(tmp_path):
    cfg = ExperimentConfig(matches_per_opponent=20, validation_matches=60,
                           brkga=BRKGAConfig(population_size=12, generations=2, seed=1))
    result = run_experiment(cfg)
    validation = result.validation
    assert not result.resampled  # motor substituto é determinístico
    assert validation.sideboard.n_matches == validation.baseline.n_matches == 60
    # Mesmos Games 1 para os dois; a linha de base não troca nenhuma carta.
    assert validation.sideboard.winrate_md1 == validation.baseline.winrate_md1
    assert all(m.cards_in == "—" for m in validation.baseline.matchups)
    low, high = validation.gain_interval()
    assert low < validation.gain < high and validation.gain_stderr > 0

    data = json.loads(export_json(result, tmp_path / "r.json").read_text(encoding="utf-8"))
    assert data["validation"]["n_matches"] == 60
    assert data["validation"]["gain"] == pytest.approx(validation.gain)


def test_search_result_is_handed_over_before_validation(tmp_path, monkeypatch):
    """Se a validação falhar, o resultado da busca já foi entregue (e pode ser gravado)."""
    def fail(self, sideboards, n_matches):
        raise ForgeError("falha na validação")

    monkeypatch.setattr(FitnessEvaluator, "validate", fail)
    cfg = ExperimentConfig(matches_per_opponent=20, brkga=BRKGAConfig(population_size=12, generations=2, seed=1))
    path = tmp_path / "parcial.json"
    with pytest.raises(ForgeError):
        run_experiment(cfg, on_search_done=lambda partial: export_json(partial, path))

    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["validation"] is None
    assert len(data["final"]["sideboard"]) == 15 and len(data["history"]) == 3
    assert data["swap_analysis"]["opponents"]


def test_measure_game1_reports_wins_per_opponent(db):
    results = measure_game1(ExperimentConfig(), 40)
    assert [opp.name for opp, _ in results] == [o.name for o in db.opponents]
    assert all(0 <= wins <= 40 for _, wins in results)


# --------------------------------------------------------------------------- #
# Arquivo de jogos
# --------------------------------------------------------------------------- #

def test_game_archive_round_trip_and_context(tmp_path):
    key = ("Bogles", (("Wear // Tear",), ("Solitude",)), 7, 2, True)
    first = GameArchive(tmp_path, "contexto-a")
    assert first.previous == [] and first.save({}) is None  # execução sem jogos não cria arquivo
    first.save({key: True})
    first.save({key: True, (*key[:3], 3, False): False})  # regrava o arquivo da própria execução

    second = GameArchive(tmp_path, "contexto-a")
    assert dict(second.previous) == {key: True, (*key[:3], 3, False): False}
    assert second.path != first.path
    # Motor ruidoso: o mesmo deck e a mesma série, jogados de novo, são um jogo novo.
    assert len(second.merged({key: False}, deterministic=False)) == 3
    # Motor determinístico: a mesma chave é o mesmo jogo, contado uma vez.
    assert len(second.merged({key: True}, deterministic=True)) == 2

    other = GameArchive(tmp_path, "contexto-b")  # outro motor ou outros decks
    assert other.previous == [] and other.skipped_runs == 1


def test_archived_games_feed_the_swap_analysis_but_not_the_results(tmp_path):
    def run(games_dir, plan_seed):
        cfg = ExperimentConfig(matches_per_opponent=20, plan_seed=plan_seed, games_dir=games_dir,
                               brkga=BRKGAConfig(population_size=8, generations=1, seed=1))
        return run_experiment(cfg)

    first = run(str(tmp_path), plan_seed=1)
    second = run(str(tmp_path), plan_seed=2)
    alone = run(None, plan_seed=2)
    assert first.games_archived == 0 and second.games_archived == first.games_simulated
    # Os jogos arquivados não entram no cache: busca e validação saem iguais às de uma execução isolada.
    assert second.best.report == alone.best.report and second.validation == alone.validation
    assert second.games_simulated == alone.games_simulated
    # ...mas a análise das trocas usa os jogos das duas execuções.
    assert sum(a.games for a in second.swap_analysis) > sum(a.games for a in alone.swap_analysis)


# --------------------------------------------------------------------------- #
# Comparação de sideboards
# --------------------------------------------------------------------------- #

def test_paired_comparison_uses_differences_between_series(db):
    report = Bo3Simulator(db, ScriptedEngine({1: True, 2: True, 3: True}), matches_per_opponent=8).evaluate(())

    def candidate(name, wins):
        return Candidate(name, (), report, {opp.name: wins for opp in db.opponents})

    wins = [True, False, True, True, False, True, False, True]
    same = paired_comparison(candidate("a", wins), candidate("b", wins))
    assert same.gain == 0 and same.stderr == 0 and not same.significant
    # Diferenças por série: uma vitória a mais em 8, igual em todos os oponentes.
    better = [True, True, True, True, False, True, False, True]
    c = paired_comparison(candidate("a", better), candidate("b", wins), block=2)
    shares = np.array([o.meta_share for o in db.opponents]) / sum(o.meta_share for o in db.opponents)
    sums = np.array([1.0, 0.0, 0.0, 0.0])  # soma das diferenças em cada bloco de 2 séries
    per_opponent = np.sum((sums - sums.mean()) ** 2) * 4 / 3 / 8 ** 2
    assert c.stderr == pytest.approx(np.sqrt(np.sum(shares ** 2) * per_opponent))
    assert np.isnan(paired_comparison(candidate("a", better), candidate("b", wins), block=8).stderr)


def test_compare_sideboards_shares_game1_and_exports(db, tmp_path):
    rng = np.random.default_rng(0)
    sideboards = {"referencia": load_reference_sideboard(), "aleatorio_1": random_sideboard(db, rng)}
    assert len(sideboard_from_names(db, sideboards["aleatorio_1"])) == 15
    with pytest.raises(ValueError):
        sideboard_from_names(db, ["Lightning Bolt"] * 15)   # fora do pool
    with pytest.raises(ValueError):
        sideboard_from_names(db, ["Blood Moon"] * 15)       # cópias demais

    cfg = ExperimentConfig(validation_matches=40, games_dir=str(tmp_path / "jogos"))
    result = compare_sideboards(cfg, sideboards)
    assert [c.name for c in result.candidates] == ["referencia", "aleatorio_1", NO_SIDEBOARD]
    assert len({c.report.winrate_md1 for c in result.candidates}) == 1  # mesmos Games 1
    assert all(c.report.n_matches == 40 for c in result.candidates)
    assert [(c.candidate, c.reference) for c in result.comparisons] == [
        ("referencia", NO_SIDEBOARD), ("aleatorio_1", NO_SIDEBOARD), ("referencia", "aleatorio_1")]
    for c in result.comparisons:
        by_name = {x.name: x.report.winrate_bo3 for x in result.candidates}
        assert c.gain == pytest.approx(by_name[c.candidate] - by_name[c.reference])
    assert result.swap_analysis and list((tmp_path / "jogos").glob("*.json"))

    data = json.loads(export_comparison_json(result, tmp_path / "c.json").read_text(encoding="utf-8"))
    assert [c["name"] for c in data["candidates"]] == ["referencia", "aleatorio_1", NO_SIDEBOARD]
    assert len(data["comparisons"]) == 3 and sum(n for _, n in data["candidates"][0]["cards"]) == 15

    from openpyxl import load_workbook

    wb = load_workbook(export_comparison_xlsx(result, tmp_path / "c.xlsx"))
    assert wb.sheetnames == ["Candidatos", "Comparacoes", "Matchups", "Efeito_Trocas", "Resumo"]
    assert wb["Matchups"].max_row == 1 + 3 * len(db.opponents)


def test_compare_command_reads_the_sideboard_of_a_previous_result(tmp_path, capsys):
    import importlib.util
    from pathlib import Path

    spec = importlib.util.spec_from_file_location("tccmagic_main", Path(__file__).resolve().parent.parent / "main.py")
    main = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(main)

    cfg = ExperimentConfig(matches_per_opponent=10, brkga=BRKGAConfig(population_size=6, generations=1, seed=1))
    previous = run_experiment(cfg)
    source = export_json(previous, tmp_path / "resultado.json")

    main.run(main.parse_args(["--compare", str(source), "--validation-matches", "20", "--random-sideboards", "1",
                              "--out", str(tmp_path / "saida"), "--games-dir", ""]))
    data = json.loads((tmp_path / "saida" / "comparacao.json").read_text(encoding="utf-8"))
    assert [c["name"] for c in data["candidates"]] == ["otimizado", "referencia", "aleatorio_1", NO_SIDEBOARD]
    assert dict(map(tuple, data["candidates"][0]["cards"])) == dict(previous.best.decoded.card_counts())
    assert data["config"]["games_dir"] is None and data["games_from_archive"] == 0
    assert "DIFERENÇAS DE WINRATE BO3" in capsys.readouterr().out

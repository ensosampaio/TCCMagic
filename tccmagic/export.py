"""Exportação do histórico de gerações, pesos finais e sideboard para JSON e XLSX."""

from __future__ import annotations

import json
from dataclasses import asdict
from pathlib import Path

from tccmagic.analysis import SwapEffect, metagame_effects
from tccmagic.attributes import ATTRIBUTES
from tccmagic.pipeline import ComparisonResult, ExperimentResult
from tccmagic.simulation import EvaluationReport


def _report_dict(report: EvaluationReport) -> dict:
    return {
        "winrate_md1": report.winrate_md1,
        "winrate_bo3": report.winrate_bo3,
        "winrate_postboard": report.winrate_postboard,
        "delta_winrate": report.delta_winrate,
        "n_matches": report.n_matches,
        "matchups": [
            {
                **asdict(m),
                "cards_in": m.cards_in,
                "cards_out": m.cards_out,
                "winrate_md1": m.winrate_md1,
                "winrate_bo3": m.winrate_bo3,
                "winrate_postboard": m.winrate_postboard,
                "delta_winrate": m.delta_winrate,
            }
            for m in report.matchups
        ],
    }


def _validation_dict(result: ExperimentResult) -> dict | None:
    validation = result.validation
    if validation is None:  # resultado parcial: a validação ainda não foi jogada
        return None
    low, high = validation.gain_interval()
    return {
        "n_matches": validation.n_matches,
        "sideboard": {**_report_dict(validation.sideboard), "stderr_bo3": validation.sideboard.stderr_bo3},
        "baseline_no_sideboard": {**_report_dict(validation.baseline), "stderr_bo3": validation.baseline.stderr_bo3},
        "gain": validation.gain,
        "gain_stderr": validation.gain_stderr,
        "gain_ci95": [low, high],
        "significant": validation.significant,
    }


def _effect_dict(e: SwapEffect) -> dict:
    low, high = e.ci95
    return {"card": e.card, "role": e.role, "games": e.games, "winrate": e.winrate, "effect": e.effect,
            "stderr": e.stderr, "ci95": [low, high], "significant": e.significant}


def _swap_analysis_dict(result: ExperimentResult | ComparisonResult) -> dict:
    return {
        "method": "probabilidade linear por oponente; entra = ganho por cópia no lugar de uma flex média; "
                  "sai = ganho de tirar a carta em vez de uma flex média",
        "opponents": [
            {"opponent": a.opponent, "archetype": a.archetype, "meta_share": a.meta_share, "games": a.games,
             "winrate_postboard": a.winrate, "no_swap_winrate": a.no_swap_winrate,
             "effects": [_effect_dict(e) for e in (*a.role("entra"), *a.role("sai"))]}
            for a in result.swap_analysis
        ],
        "metagame": [_effect_dict(e) for e in metagame_effects(result.swap_analysis)],
    }


def result_to_dict(result: ExperimentResult) -> dict:
    best = result.best
    return {
        "config": result.config_dict(),
        "maindeck": result.database.maindeck.name,
        "elapsed_seconds": result.elapsed_seconds,
        "stopped_early": result.stopped_early,
        "resampled": result.resampled,
        "games_simulated": result.games_simulated,
        "games_reused": result.games_reused,
        "games_from_archive": result.games_archived,
        "final": {
            "fitness": best.fitness,
            "keys": dict(zip(ATTRIBUTES, result.best_keys)),
            "weights": best.decoded.weights_by_attribute(),
            "sideboard": [
                {"rank": i, "card": slot.card.name, "copy": slot.copy_index, "score": score}
                for i, (slot, score) in enumerate(zip(best.decoded.slots, best.decoded.scores), start=1)
            ],
            "evaluation": _report_dict(best.report),
        },
        "validation": _validation_dict(result),
        "swap_analysis": _swap_analysis_dict(result),
        "history": [asdict(rec) for rec in result.history],
    }


def export_json(result: ExperimentResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(result_to_dict(result), fh, ensure_ascii=False, indent=2)
    return path


MATCHUP_HEADER = ["oponente", "arquetipo", "meta_share", "series", "winrate_md1", "winrate_bo3",
                  "winrate_pos_side", "delta_winrate", "entram (copias por serie)", "saem (copias por serie)"]
EFFECT_HEADER = ["oponente", "carta", "papel", "jogos", "winrate_bruta", "efeito_por_copia", "erro_padrao",
                 "ic95_inferior", "ic95_superior", "significativo"]


def _matchup_row(m) -> list:
    return [m.opponent, m.archetype, m.meta_share, m.n_matches, m.winrate_md1, m.winrate_bo3,
            m.winrate_postboard, m.delta_winrate, m.cards_in, m.cards_out]


def _effect_rows(result: ExperimentResult | ComparisonResult) -> list[list]:
    effects = [*(e for a in result.swap_analysis for e in (*a.role("entra"), *a.role("sai"))),
               *metagame_effects(result.swap_analysis)]
    return [[e.opponent, e.card, e.role, e.games, e.winrate, e.effect, e.stderr, *e.ci95, e.significant]
            for e in effects]


def _workbook():
    """Planilha nova e a função que acrescenta uma aba com cabeçalho em negrito."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise ImportError("Exportação .xlsx requer 'openpyxl' (pip install openpyxl).") from exc

    wb = Workbook()

    def sheet(title: str, header: list[str], rows: list[list], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for column in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(60, width + 2)
        return ws

    return wb, sheet


def export_xlsx(result: ExperimentResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if result.validation is None:
        raise ValueError("Resultado parcial (sem validação): exporte com export_json.")
    wb, sheet = _workbook()
    best = result.best

    sheet(
        "Geracoes",
        ["geracao", "melhor_fitness", "fitness_medio", "fitness_desvio", "winrate_md1", "winrate_bo3",
         "winrate_pos_side", "delta_winrate", "series_do_melhor", "sideboards_simulados",
         "jogos_simulados", "jogos_reaproveitados", "tempo_s", *[f"w_{a}" for a in ATTRIBUTES]],
        [[r.generation, r.best_fitness, r.mean_fitness, r.std_fitness, r.winrate_md1, r.winrate_bo3,
          r.winrate_postboard, r.delta_winrate, r.best_n_matches, r.simulated_sideboards, r.games_simulated,
          r.games_reused, r.elapsed_seconds, *[r.weights[a] for a in ATTRIBUTES]] for r in result.history],
        first=True,
    )
    sheet(
        "Pesos_Finais",
        ["atributo", "chave", "peso"],
        [[a, k, w] for a, k, w in zip(ATTRIBUTES, result.best_keys, best.decoded.weights)],
    )
    sheet(
        "Sideboard",
        ["rank", "carta", "copia", "score"],
        [[i, s.card.name, s.copy_index, sc]
         for i, (s, sc) in enumerate(zip(best.decoded.slots, best.decoded.scores), start=1)],
    )
    sheet("Matchups", MATCHUP_HEADER, [_matchup_row(m) for m in best.report.matchups])
    validation = result.validation
    sheet(
        "Validacao",
        ["deck", *MATCHUP_HEADER],
        [[label, *_matchup_row(m)]
         for label, report in (("com_sideboard", validation.sideboard), ("sem_sideboard", validation.baseline))
         for m in report.matchups],
    )
    sheet("Efeito_Trocas", EFFECT_HEADER, _effect_rows(result))
    low, high = validation.gain_interval()
    sheet(
        "Resumo",
        ["metrica", "valor"],
        [
            ["maindeck", result.database.maindeck.name],
            ["motor", result.config.engine],
            ["plano_de_troca", result.config.plan],
            ["jogos_pos_side_simulados", result.games_simulated],
            ["jogos_pos_side_reaproveitados_do_cache", result.games_reused],
            ["jogos_de_execucoes_anteriores_na_analise_de_trocas", result.games_archived],
            ["geracoes_executadas", len(result.history) - 1],
            ["fitness_final_na_busca (winrate_bo3, otimista)", best.fitness],
            ["series_por_oponente_na_busca", best.report.n_matches],
            ["validacao_series_por_oponente", validation.n_matches],
            ["validacao_winrate_md1", validation.sideboard.winrate_md1],
            ["validacao_winrate_bo3", validation.sideboard.winrate_bo3],
            ["validacao_winrate_bo3_erro_padrao", validation.sideboard.stderr_bo3],
            ["validacao_delta_winrate (bo3 - md1)", validation.sideboard.delta_winrate],
            ["validacao_winrate_bo3_sem_sideboard", validation.baseline.winrate_bo3],
            ["ganho_do_sideboard (bo3 - bo3_sem_side)", validation.gain],
            ["ganho_ic95_inferior", low],
            ["ganho_ic95_superior", high],
            ["ganho_significativo (ic95 exclui zero)", validation.significant],
            ["tempo_total_s", result.elapsed_seconds],
        ],
    )
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# Comparação de sideboards (python main.py --compare ...)
# --------------------------------------------------------------------------- #

def comparison_to_dict(result: ComparisonResult) -> dict:
    return {
        "config": result.config_dict(),
        "maindeck": result.database.maindeck.name,
        "elapsed_seconds": result.elapsed_seconds,
        "n_matches": result.n_matches,
        "games_simulated": result.games_simulated,
        "games_reused": result.games_reused,
        "games_from_archive": result.games_archived,
        "candidates": [
            {"name": c.name, "cards": [list(pair) for pair in c.cards],
             **_report_dict(c.report), "stderr_bo3": c.report.stderr_bo3}
            for c in result.candidates
        ],
        "comparisons": [
            {"candidate": c.candidate, "reference": c.reference, "gain": c.gain, "stderr": c.stderr,
             "ci95": list(c.interval()), "significant": c.significant}
            for c in result.comparisons
        ],
        "swap_analysis": _swap_analysis_dict(result),
    }


def export_comparison_json(result: ComparisonResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(comparison_to_dict(result), fh, ensure_ascii=False, indent=2)
    return path


def export_comparison_xlsx(result: ComparisonResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb, sheet = _workbook()
    sheet(
        "Candidatos",
        ["sideboard", "cartas", "series", "winrate_md1", "winrate_bo3", "erro_padrao_bo3", "winrate_pos_side"],
        [[c.name, ", ".join(f"{n} {card}" for card, n in c.cards) or "—", c.report.n_matches,
          c.report.winrate_md1, c.report.winrate_bo3, c.report.stderr_bo3, c.report.winrate_postboard]
         for c in result.candidates],
        first=True,
    )
    sheet(
        "Comparacoes",
        ["sideboard", "comparado_com", "ganho_bo3", "erro_padrao_pareado", "ic95_inferior", "ic95_superior",
         "significativo"],
        [[c.candidate, c.reference, c.gain, c.stderr, *c.interval(), c.significant] for c in result.comparisons],
    )
    sheet(
        "Matchups",
        ["sideboard", *MATCHUP_HEADER],
        [[c.name, *_matchup_row(m)] for c in result.candidates for m in c.report.matchups],
    )
    sheet("Efeito_Trocas", EFFECT_HEADER, _effect_rows(result))
    sheet(
        "Resumo",
        ["metrica", "valor"],
        [
            ["maindeck", result.database.maindeck.name],
            ["motor", result.config.engine],
            ["plano_de_troca", result.config.plan],
            ["series_por_oponente", result.n_matches],
            ["jogos_pos_side_simulados", result.games_simulated],
            ["jogos_pos_side_reaproveitados_do_cache", result.games_reused],
            ["jogos_de_execucoes_anteriores_na_analise_de_trocas", result.games_archived],
            ["tempo_total_s", result.elapsed_seconds],
        ],
    )
    wb.save(path)
    return path
